#!/usr/bin/env python3
"""
Converter: the German spreadsheet -> content/de/de-a1-b2/items.json

Run with `pnpm xlsx:convert`. After this the app only ever reads the JSON files;
the spreadsheet is kept in data-source/ purely for traceability. course.json is
maintained by hand and is never written here.

Only rows that actually carry a German sentence are emitted; empty placeholder
rows are reported and skipped, never invented.

The sheets are read by column position, so the header row is checked first: a
reordered or renamed column would otherwise be imported silently into the wrong
field.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Cần openpyxl: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
COURSE_ID = "de-a1-b2"
OUT_DIR = ROOT / "content" / "de" / COURSE_ID

# Candidate locations for the source workbook, first match wins.
SOURCE_CANDIDATES = [
    ROOT / "data-source" / "1800 Completed.xlsx",
    Path.home() / "Downloads" / "1800 Completed.xlsx",
    ROOT / "data-source" / "Customized Learning Deustch.xlsx",
    Path.home() / "Downloads" / "Customized Learning Deustch.xlsx",
]

# Column positions the reader depends on, as (index, expected header fragment).
# Matched case-insensitively on a prefix so a cosmetic rewording of a header does
# not break the import, while a moved column does.
SENTENCE_COLUMNS = [
    (0, "id"),
    (3, "cấp độ"),
    (4, "chủ đề"),
    (5, "loại câu"),
    (6, "điểm ngữ pháp"),
    (7, "câu tiếng đức"),
    (8, "phiên âm"),
    (9, "nghĩa tiếng việt"),
    (10, "từ mới"),
]

VOCAB_COLUMNS = [
    (1, "từ"),
    (2, "nghĩa tiếng việt"),
    (3, "loại từ"),
    (4, "giống"),
    (5, "id câu"),
]

# --- Vietnamese spreadsheet labels -> canonical slugs ------------------------
# Canonical values are language-neutral; Vietnamese labels live in src/messages/vi.json.
SENTENCE_TYPE_MAP = {
    "Câu trần thuật": "declarative",
    "Câu hỏi": "question",
    "Câu hỏi Ja/Nein": "yes_no_question",
    "Câu phủ định": "negative",
    "Mệnh lệnh": "imperative",
    "Câu cảm thán": "exclamation",
    "Câu ghép": "compound",
    "Câu phụ": "subordinate",
    "Câu trả lời": "answer",
    "Câu bị động": "passive",
    "Liệt kê": "enumeration",
    # Added with the completed B1-B2 half of the workbook. "Câu phụ" and
    # "Mệnh đề phụ" are kept apart because the source keeps them apart; merging
    # them would be editing the data, not importing it.
    "Mệnh đề phụ": "subordinate_clause",
    "Mệnh đề quan hệ": "relative_clause",
    "Mệnh đề nguyên mẫu": "infinitive_clause",
    "Thành ngữ": "idiom",
    "Khẩu ngữ": "colloquial",
    "Văn viết": "formal_written",
    "Lưu ý lỗi": "common_mistake",
    "Mẫu câu thi": "exam_pattern",
}

PART_OF_SPEECH_MAP = {
    "Danh từ": "noun",
    "Danh từ (không rõ giống)": "noun",
    "Động từ": "verb",
    "Động từ phản thân": "reflexive_verb",
    "Cụm từ / Thành ngữ": "phrase",
    "Tính từ / Trạng từ / Khác": "other",
}

GENDER_MAP = {
    "Đực (der)": "der",
    "Cái (die)": "die",
    "Trung (das)": "das",
}


def norm(value: object) -> str:
    """Trim, collapse inner whitespace, normalise unicode. Empty for None."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    return re.sub(r"\s+", " ", text).strip()


def order_index_from_key(stable_key: str) -> int:
    match = re.match(r"^[A-Z](\d+)$", stable_key)
    if not match:
        raise ValueError(f"stableKey không hợp lệ: {stable_key!r}")
    return int(match.group(1))


def check_columns(ws, columns: list[tuple[int, str]], sheet: str) -> list[str]:
    """Confirm the header row still matches the positions the reader assumes."""
    header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), ())
    wrong: list[str] = []
    for index, expected in columns:
        actual = norm(header[index]).lower() if index < len(header) else ""
        if not actual.startswith(expected):
            wrong.append(
                f"{sheet} cột {index + 1}: cần tiêu đề bắt đầu bằng {expected!r}, đang là {actual!r}"
            )
    return wrong


def find_source() -> Path:
    for candidate in SOURCE_CANDIDATES:
        if candidate.is_file():
            return candidate
    locations = "\n  ".join(str(c) for c in SOURCE_CANDIDATES)
    sys.exit(f"Không tìm thấy file Excel nguồn. Đã thử:\n  {locations}")


def main() -> int:
    source = find_source()
    print(f"Đọc: {source}")
    wb = openpyxl.load_workbook(source, data_only=True)

    for required in ("SENTENCES", "VOCAB"):
        if required not in wb.sheetnames:
            sys.exit(f"File Excel thiếu sheet {required!r}.")

    # Stop before reading a single row if the columns have moved: importing the
    # wrong column into the right field is far worse than failing here.
    layout = check_columns(wb["SENTENCES"], SENTENCE_COLUMNS, "SENTENCES") + check_columns(
        wb["VOCAB"], VOCAB_COLUMNS, "VOCAB"
    )
    if layout:
        print("✗ Bố cục cột không khớp với những gì script mong đợi:")
        for problem in layout:
            print(f"  - {problem}")
        sys.exit("Sửa file Excel, hoặc cập nhật SENTENCE_COLUMNS / VOCAB_COLUMNS trong script này.")

    # --- glossary, keyed by sentence id -----------------------------------
    glossary_by_item: dict[str, list[dict[str, object]]] = defaultdict(list)
    vocab_problems: list[str] = []
    vocab_rows = 0
    for row_no, row in enumerate(wb["VOCAB"].iter_rows(min_row=3, values_only=True), start=3):
        term, meaning, pos_label, gender_label, item_id = (
            norm(row[1]),
            norm(row[2]),
            norm(row[3]),
            norm(row[4]),
            norm(row[5]),
        )
        if not term:
            continue
        vocab_rows += 1
        if not meaning:
            vocab_problems.append(f"VOCAB dòng {row_no}: từ {term!r} thiếu nghĩa.")
            continue
        if not item_id:
            vocab_problems.append(f"VOCAB dòng {row_no}: từ {term!r} không gắn với ID câu nào.")
            continue
        pos = PART_OF_SPEECH_MAP.get(pos_label)
        if pos is None:
            vocab_problems.append(f"VOCAB dòng {row_no}: loại từ lạ {pos_label!r} (từ {term!r}).")
            pos = "other"
        entry: dict[str, object] = {"term": term, "meaning": meaning, "partOfSpeech": pos, "fields": {}}
        if gender_label:
            gender = GENDER_MAP.get(gender_label)
            if gender is None:
                vocab_problems.append(f"VOCAB dòng {row_no}: giống lạ {gender_label!r} (từ {term!r}).")
            else:
                entry["fields"] = {"gender": gender}
        glossary_by_item[item_id].append(entry)

    # --- items ------------------------------------------------------------
    items: list[dict[str, object]] = []
    skipped_empty: list[str] = []
    problems: list[str] = list(vocab_problems)
    levels_seen: Counter[str] = Counter()
    topics_order: list[str] = []

    for row_no, row in enumerate(wb["SENTENCES"].iter_rows(min_row=3, values_only=True), start=3):
        stable_key = norm(row[0])
        if not stable_key:
            continue
        text = norm(row[7])
        if not text:
            # Placeholder row: id and planned date exist, content does not.
            skipped_empty.append(stable_key)
            continue

        level = norm(row[3])
        topic = norm(row[4])
        type_label = norm(row[5])
        grammar_note = norm(row[6])
        transliteration = norm(row[8])
        meaning = norm(row[9])

        if not level:
            problems.append(f"SENTENCES dòng {row_no} ({stable_key}): thiếu cấp độ.")
        if not topic:
            problems.append(f"SENTENCES dòng {row_no} ({stable_key}): thiếu chủ đề.")
        if not meaning:
            problems.append(f"SENTENCES dòng {row_no} ({stable_key}): thiếu nghĩa tiếng Việt.")

        fields: dict[str, str] = {"text": text, "meaning": meaning}
        if transliteration:
            fields["transliteration"] = transliteration
        if grammar_note:
            fields["grammarNote"] = grammar_note
        if type_label:
            sentence_type = SENTENCE_TYPE_MAP.get(type_label)
            if sentence_type is None:
                problems.append(f"SENTENCES dòng {row_no} ({stable_key}): loại câu lạ {type_label!r}.")
            else:
                fields["sentenceType"] = sentence_type

        # Cross-check the free-text "Từ mới trong câu" column against VOCAB.
        inline_words = [w for w in (norm(row[10]) or "").split(";") if w.strip()]
        glossary = glossary_by_item.get(stable_key, [])
        if inline_words and len(inline_words) != len(glossary):
            problems.append(
                f"SENTENCES dòng {row_no} ({stable_key}): cột 'Từ mới' có {len(inline_words)} từ "
                f"nhưng VOCAB có {len(glossary)} từ."
            )

        levels_seen[level] += 1
        if topic and topic not in topics_order:
            topics_order.append(topic)

        items.append(
            {
                "stableKey": stable_key,
                "kind": "sentence",
                "level": level,
                "topic": topic,
                "orderIndex": order_index_from_key(stable_key),
                "fields": fields,
                "glossary": glossary,
            }
        )

    items.sort(key=lambda it: it["orderIndex"])

    orphan_ids = sorted(set(glossary_by_item) - {it["stableKey"] for it in items})
    for orphan in orphan_ids:
        problems.append(f"VOCAB trỏ tới ID câu {orphan} nhưng câu đó không có nội dung.")

    # --- write ------------------------------------------------------------
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    items_path = OUT_DIR / "items.json"
    payload = {"courseId": COURSE_ID, "schemaVersion": 1, "items": items}
    items_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    glossary_total = sum(len(it["glossary"]) for it in items)
    distinct_terms = len({(g["term"], g["meaning"]) for it in items for g in it["glossary"]})

    print(f"\n✓ Ghi {items_path.relative_to(ROOT)}")
    print(f"  Mục có nội dung : {len(items)}")
    print(f"  Bỏ qua (rỗng)   : {len(skipped_empty)}"
          + (f"  [{skipped_empty[0]}..{skipped_empty[-1]}]" if skipped_empty else ""))
    print(f"  Cấp độ          : {dict(sorted(levels_seen.items()))}")
    print(f"  Chủ đề          : {len(topics_order)}")
    print(f"  Dòng VOCAB      : {vocab_rows} -> {glossary_total} liên kết, {distinct_terms} từ riêng biệt")

    if problems:
        print(f"\n⚠ {len(problems)} vấn đề cần xem:")
        for p in problems[:30]:
            print(f"  - {p}")
        if len(problems) > 30:
            print(f"  ... và {len(problems) - 30} vấn đề khác")
    else:
        print("\n✓ Không phát hiện vấn đề nào trong dữ liệu nguồn.")

    print("\nBước tiếp: pnpm content:validate")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
